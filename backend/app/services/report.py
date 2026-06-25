"""
WHAT: Assembles a professional, period-scoped financial report — the chief-of-staff
      brief in document form. Powers GET /reports/financial (+ a CSV appendix).
WHY: A useful report leads with judgment and frames every figure as change-over-period,
      not a transaction dump. It reuses the same conversion + ledger logic as the net-worth
      summary so the document agrees with the app. All text is deterministic (reports must
      be correct), built from the numbers — no LLM in the loop.
BREAKS IF REMOVED: /reports endpoints have nothing to assemble.
"""

import csv
import io
import logging
import uuid
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.asset import Asset
from app.models.liability import Liability
from app.models.networth_snapshot import NetworthSnapshot
from app.models.receivable import Receivable
from app.models.transaction import Transaction
from app.services.currency import convert

logger = logging.getLogger(__name__)

LIQUID_TYPES = {"cash", "bank_account"}

# Category slug → label, per language (mirrors the rest of the app).
_CAT = {
    "tr": {"market": "Market", "restoran": "Restoran", "ulasim": "Ulaşım", "eglence": "Eğlence",
           "saglik": "Sağlık", "fatura": "Fatura", "giyim": "Giyim", "nakit_atm": "Nakit/ATM",
           "transfer": "Transfer", "faiz": "Faiz", "iade": "İade", "vergi": "Vergi", "teknoloji": "Teknoloji",
           "diger": "Diğer", "egitim": "Eğitim"},
    "en": {"market": "Groceries", "restoran": "Dining", "ulasim": "Transport", "eglence": "Entertainment",
           "saglik": "Health", "fatura": "Bills", "giyim": "Clothing", "nakit_atm": "Cash/ATM",
           "transfer": "Transfer", "faiz": "Interest", "iade": "Refund", "vergi": "Tax", "teknoloji": "Technology",
           "diger": "Other", "egitim": "Education"},
}

_ASSET_TYPE = {
    "tr": {"cash": "Nakit", "bank_account": "Banka Hesabı", "crypto": "Kripto", "stock": "Hisse",
           "fund": "Fon", "gold": "Altın", "foreign_currency": "Döviz", "commodity": "Emtia",
           "real_estate": "Gayrimenkul", "vehicle": "Araç", "bond": "Tahvil", "bes": "BES",
           "pension": "Emeklilik", "life_insurance": "Hayat Sigortası", "art_collectible": "Sanat/Koleksiyon",
           "jewelry": "Mücevher", "business_ownership": "İşletme", "other_asset": "Diğer"},
    "en": {"cash": "Cash", "bank_account": "Bank account", "crypto": "Crypto", "stock": "Stocks",
           "fund": "Funds", "gold": "Gold", "foreign_currency": "Foreign currency", "commodity": "Commodity",
           "real_estate": "Real estate", "vehicle": "Vehicle", "bond": "Bonds", "bes": "Pension (BES)",
           "pension": "Pension", "life_insurance": "Life insurance", "art_collectible": "Art/Collectible",
           "jewelry": "Jewelry", "business_ownership": "Business", "other_asset": "Other"},
}


def _cat_label(slug: str, lang: str) -> str:
    return _CAT.get(lang, _CAT["en"]).get(slug, slug)


def _type_label(slug: str, lang: str) -> str:
    return _ASSET_TYPE.get(lang, _ASSET_TYPE["en"]).get(slug, slug)


def _money(n: float, ccy: str) -> str:
    return f"{n:,.0f} {ccy}"


# ── period resolution ────────────────────────────────────────────────────────

def resolve_period(key: str, today: date | None = None) -> tuple[date | None, date, str]:
    today = today or date.today()
    if key == "this_month":
        return date(today.year, today.month, 1), today, "this_month"
    if key == "last_month":
        first_this = date(today.year, today.month, 1)
        end = first_this - timedelta(days=1)
        return date(end.year, end.month, 1), end, "last_month"
    if key == "quarter":
        q = (today.month - 1) // 3
        return date(today.year, q * 3 + 1, 1), today, "quarter"
    if key == "ytd":
        return date(today.year, 1, 1), today, "ytd"
    if key == "last_30":
        return today - timedelta(days=30), today, "last_30"
    return None, today, "all"


def _period_label(key: str, start: date | None, end: date, lang: str) -> str:
    if start is None:
        return ("Tüm zamanlar" if lang == "tr" else "All time") + f" – {end.isoformat()}"
    return f"{start.isoformat()} → {end.isoformat()}"


async def _factor(from_ccy: str, to_ccy: str, cache: dict[str, float]) -> float:
    fc = (from_ccy or to_ccy).upper()
    if fc == to_ccy.upper():
        return 1.0
    if fc not in cache:
        try:
            cache[fc] = float(await convert(1.0, fc, to_ccy))
        except Exception:  # noqa: BLE001
            cache[fc] = 1.0
    return cache[fc]


# ── report assembly ──────────────────────────────────────────────────────────

async def build_report(
    user_id: uuid.UUID, session: AsyncSession, period_key: str, ccy: str, lang: str,
) -> dict:
    ccy = ccy.upper()
    lang = lang if lang in ("tr", "en") else "en"
    start, end, period_key = resolve_period(period_key)
    cache: dict[str, float] = {}

    # ── ledger (current position) ──
    assets = (await session.execute(select(Asset).where(Asset.user_id == user_id))).scalars().all()
    liabs = (await session.execute(select(Liability).where(Liability.user_id == user_id))).scalars().all()
    recvs = (await session.execute(
        select(Receivable).where(Receivable.user_id == user_id, Receivable.status == "pending")
    )).scalars().all()

    total_assets = 0.0
    liquid = 0.0
    by_type: dict[str, float] = {}
    by_currency: dict[str, float] = {}
    asset_rows: list[dict] = []
    for a in assets:
        f = await _factor(a.currency, ccy, cache)
        val = float(a.current_value) * f
        total_assets += val
        by_type[a.asset_type] = by_type.get(a.asset_type, 0.0) + val
        by_currency[a.currency] = by_currency.get(a.currency, 0.0) + val
        if a.asset_type in LIQUID_TYPES:
            liquid += val
        asset_rows.append({"name": a.name, "type": _type_label(a.asset_type, lang), "value": round(val)})
    asset_rows.sort(key=lambda r: r["value"], reverse=True)

    total_liab = 0.0
    liab_rows: list[dict] = []
    for l in liabs:
        f = await _factor(l.currency, ccy, cache)
        rem = float(l.remaining_amount) * f
        total_liab += rem
        liab_rows.append({
            "name": l.name, "remaining": round(rem),
            "monthly_payment": round(float(l.monthly_payment or 0) * f),
            "rate": float(l.interest_rate) if l.interest_rate is not None else None,
        })
    liab_rows.sort(key=lambda r: r["remaining"], reverse=True)

    recv_total = 0.0
    recv_rows: list[dict] = []
    for r in recvs:
        f = await _factor(r.currency, ccy, cache)
        amt = float(r.amount) * f
        recv_total += amt
        recv_rows.append({"from_person": r.from_person, "amount": round(amt),
                          "expected_date": r.expected_date.isoformat() if r.expected_date else None})
    recv_rows.sort(key=lambda r: r["amount"], reverse=True)

    net_worth = total_assets - total_liab

    # ── trajectory + opening/closing from snapshots ──
    snaps = (await session.execute(
        select(NetworthSnapshot).where(NetworthSnapshot.user_id == user_id)
        .order_by(NetworthSnapshot.recorded_at.asc())
    )).scalars().all()
    usd_factor = await _factor("USD", ccy, cache)
    traj = [{"date": s.recorded_at.date().isoformat(), "net_worth": round(float(s.net_worth_usd) * usd_factor)}
            for s in snaps]
    # cap to ~24 evenly-spaced points
    if len(traj) > 24:
        step = len(traj) / 24
        traj = [traj[int(i * step)] for i in range(24)] + [traj[-1]]

    opening = closing = None
    nw_estimated = True
    if snaps:
        closing = round(float(snaps[-1].net_worth_usd) * usd_factor)
        if start is not None:
            before = [s for s in snaps if s.recorded_at.date() <= start]
            ref = before[-1] if before else snaps[0]
        else:
            ref = snaps[0]
        opening = round(float(ref.net_worth_usd) * usd_factor)
        nw_estimated = len(snaps) < 2 or opening == closing
    nw_change = (closing - opening) if (opening is not None and closing is not None) else None

    # ── cash flow over the period ──
    tx_stmt = select(Transaction).where(Transaction.user_id == user_id, Transaction.transaction_date <= end)
    if start is not None:
        tx_stmt = tx_stmt.where(Transaction.transaction_date >= start)
    txs = (await session.execute(tx_stmt)).scalars().all()
    income = expense = 0.0
    cat_totals: dict[str, float] = {}
    for t in txs:
        f = await _factor(t.currency or ccy, ccy, cache)
        amt = float(t.amount) * f
        if t.transaction_type == "credit":
            income += amt
        else:
            expense += amt
            cat_totals[t.category or "diger"] = cat_totals.get(t.category or "diger", 0.0) + amt
    net_cash = income - expense
    total_exp = sum(cat_totals.values()) or 1.0
    top_categories = [
        {"name": _cat_label(slug, lang), "amount": round(v), "share": round(v / total_exp * 100, 1)}
        for slug, v in sorted(cat_totals.items(), key=lambda x: x[1], reverse=True)[:6]
    ]

    # ── allocation ──
    allocation = [
        {"name": _type_label(k, lang), "value": round(v), "share": round(v / (total_assets or 1) * 100, 1)}
        for k, v in sorted(by_type.items(), key=lambda x: x[1], reverse=True)
    ]
    currency_mix = [
        {"code": k, "value": round(v), "share": round(v / (total_assets or 1) * 100, 1)}
        for k, v in sorted(by_currency.items(), key=lambda x: x[1], reverse=True)
    ]

    # ── recommendations (deterministic, ranked, max 3) ──
    recs: list[dict] = []
    debt_ratio = total_liab / total_assets if total_assets > 0 else 0.0
    liq_ratio = liquid / total_assets if total_assets > 0 else 0.0
    top_asset_share = max((r["value"] for r in asset_rows), default=0) / (total_assets or 1)

    if net_cash < 0:
        recs.append((3, {
            "title": "Negatif nakit akışı" if lang == "tr" else "Negative cash flow",
            "detail": (f"Bu dönemde giderlerin gelirini {_money(abs(net_cash), ccy)} aştı."
                       if lang == "tr" else
                       f"Spending exceeded income by {_money(abs(net_cash), ccy)} this period."),
        }))
    if debt_ratio > 0.4:
        recs.append((2 if debt_ratio <= 0.7 else 4, {
            "title": "Yüksek borç yükü" if lang == "tr" else "High debt load",
            "detail": (f"Borçlar varlıkların %{debt_ratio*100:.0f}'ini oluşturuyor."
                       if lang == "tr" else
                       f"Debt is {debt_ratio*100:.0f}% of your assets."),
        }))
    if total_assets > 0 and liq_ratio < 0.1:
        recs.append((2, {
            "title": "Düşük nakit tamponu" if lang == "tr" else "Thin cash buffer",
            "detail": (f"Likit varlıklar toplamın yalnızca %{liq_ratio*100:.0f}'i."
                       if lang == "tr" else
                       f"Liquid assets are only {liq_ratio*100:.0f}% of the total."),
        }))
    if top_asset_share > 0.4 and asset_rows:
        recs.append((1, {
            "title": "Tek varlıkta yoğunlaşma" if lang == "tr" else "Concentration risk",
            "detail": (f"Varlıklarının %{top_asset_share*100:.0f}'i tek bir kalemde ({asset_rows[0]['name']})."
                       if lang == "tr" else
                       f"{top_asset_share*100:.0f}% of assets sit in one item ({asset_rows[0]['name']})."),
        }))
    recommendations = [r[1] for r in sorted(recs, key=lambda x: x[0], reverse=True)[:3]]

    # ── executive summary (deterministic verdict) ──
    summary = _exec_summary(lang, ccy, net_worth, nw_change, nw_estimated, income, expense, net_cash, recommendations)

    assumptions = (
        [
            "Tutarlar canlı döviz/varlık kurlarıyla seçtiğin para birimine çevrildi.",
            "Net değer değişimi anlık görüntülere dayanır; az veri varsa tahminîdir.",
            "Nakit akışı yalnızca yüklenen/işlenen işlemleri yansıtır.",
        ] if lang == "tr" else [
            "Amounts converted to your selected currency at live FX/asset rates.",
            "Net-worth change is based on snapshots; estimated when data is sparse.",
            "Cash flow reflects only uploaded/recorded transactions.",
        ]
    )

    return {
        "meta": {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "period_key": period_key,
            "period_label": _period_label(period_key, start, end, lang),
            "currency": ccy,
            "lang": lang,
        },
        "summary": summary,
        "net_worth": {
            "total_assets": round(total_assets),
            "total_liabilities": round(total_liab),
            "net_worth": round(net_worth),
            "pending_receivables": round(recv_total),
            "opening": opening,
            "closing": closing,
            "change": nw_change,
            "estimated": nw_estimated,
        },
        "cash_flow": {
            "income": round(income),
            "expenses": round(expense),
            "net": round(net_cash),
            "top_categories": top_categories,
        },
        "assets": asset_rows[:15],
        "liabilities": liab_rows[:15],
        "receivables": recv_rows[:10],
        "allocation": allocation,
        "currency_mix": currency_mix,
        "trajectory": traj,
        "recommendations": recommendations,
        "assumptions": assumptions,
    }


def _exec_summary(lang, ccy, net_worth, nw_change, estimated, income, expense, net_cash, recs) -> str:
    sign = lambda v: "+" if v >= 0 else "−"  # noqa: E731
    if lang == "tr":
        s = f"Net değerin {_money(net_worth, ccy)}."
        if nw_change is not None:
            s += (f" Dönem boyunca {sign(nw_change)}{_money(abs(nw_change), ccy)} "
                  f"{'arttı' if nw_change >= 0 else 'azaldı'}{' (tahminî)' if estimated else ''}.")
        s += (f" Bu dönemde {_money(income, ccy)} giriş, {_money(expense, ccy)} çıkış oldu — "
              f"{'fazla' if net_cash >= 0 else 'açık'} {_money(abs(net_cash), ccy)}.")
        if recs:
            s += f" Öncelikli konu: {recs[0]['title'].lower()}."
        return s
    s = f"Your net worth is {_money(net_worth, ccy)}."
    if nw_change is not None:
        s += (f" Over the period it {'rose' if nw_change >= 0 else 'fell'} "
              f"{sign(nw_change)}{_money(abs(nw_change), ccy)}{' (estimated)' if estimated else ''}.")
    s += (f" This period saw {_money(income, ccy)} in and {_money(expense, ccy)} out — "
          f"a {'surplus' if net_cash >= 0 else 'shortfall'} of {_money(abs(net_cash), ccy)}.")
    if recs:
        s += f" Top priority: {recs[0]['title'].lower()}."
    return s


# ── CSV appendix ─────────────────────────────────────────────────────────────

# ── Excel workbook ───────────────────────────────────────────────────────────

async def build_report_xlsx(
    user_id: uuid.UUID, session: AsyncSession, period_key: str, ccy: str, lang: str,
) -> bytes:
    """Multi-sheet .xlsx workbook of the same report — for accountants / analysts.

    Reuses build_report() for the structured figures and re-queries the raw
    transactions for the appendix sheet. openpyxl is already a dependency.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    report = await build_report(user_id, session, period_key, ccy, lang)
    ccy = ccy.upper()
    num_fmt = "#,##0"

    header_fill = PatternFill("solid", fgColor="0F5C5E")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="0F5C5E")
    label_font = Font(bold=True)

    wb = Workbook()

    def _style_header(ws, ncols: int, row: int = 1) -> None:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    def _autosize(ws, widths: list[int]) -> None:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _table(ws, headers: list[str], rows: list[list], widths: list[int],
               money_cols: tuple[int, ...] = (), pct_cols: tuple[int, ...] = ()) -> None:
        ws.append(headers)
        _style_header(ws, len(headers))
        for r in rows:
            ws.append(r)
        for ri in range(2, len(rows) + 2):
            for c in money_cols:
                ws.cell(row=ri, column=c).number_format = num_fmt
            for c in pct_cols:
                ws.cell(row=ri, column=c).number_format = '0"%"'
        _autosize(ws, widths)

    # ── Summary ──
    nw, cf, meta = report["net_worth"], report["cash_flow"], report["meta"]
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Mizan · Financial Report"
    ws["A1"].font = title_font
    ws["A2"] = meta["period_label"]
    ws["A3"] = f"Currency: {ccy}   ·   Generated: {meta['generated_at'][:10]}"
    ws["A3"].font = Font(color="64748B")
    ws["A5"] = report["summary"]
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A5:D7")

    rows = [
        ("Net worth", nw["net_worth"]),
        ("Total assets", nw["total_assets"]),
        ("Total liabilities", nw["total_liabilities"]),
        ("Pending receivables", nw["pending_receivables"]),
        ("Income", cf["income"]),
        ("Expenses", cf["expenses"]),
        ("Net cash flow", cf["net"]),
    ]
    if nw["opening"] is not None:
        rows += [("Opening net worth", nw["opening"]), ("Closing net worth", nw["closing"]), ("Net change", nw["change"])]
    start_row = 9
    ws.cell(row=start_row, column=1, value="Metric").font = header_font
    ws.cell(row=start_row, column=2, value=f"Amount ({ccy})").font = header_font
    for c in (1, 2):
        ws.cell(row=start_row, column=c).fill = header_fill
    for i, (label, val) in enumerate(rows, start=start_row + 1):
        ws.cell(row=i, column=1, value=label).font = label_font
        vc = ws.cell(row=i, column=2, value=val)
        vc.number_format = num_fmt
    _autosize(ws, [26, 20])

    # ── Holdings ──
    ws = wb.create_sheet("Holdings")
    _table(ws, ["Name", "Type", f"Value ({ccy})"],
           [[a["name"], a["type"], a["value"]] for a in report["assets"]],
           [34, 22, 18], money_cols=(3,))

    # ── Liabilities ──
    ws = wb.create_sheet("Liabilities")
    _table(ws, ["Name", f"Remaining ({ccy})", f"Monthly ({ccy})", "Rate %"],
           [[l["name"], l["remaining"], l.get("monthly_payment") or 0, l.get("rate")] for l in report["liabilities"]],
           [34, 20, 18, 10], money_cols=(2, 3))

    # ── Receivables ──
    if report["receivables"]:
        ws = wb.create_sheet("Receivables")
        _table(ws, ["From", f"Amount ({ccy})", "Expected date"],
               [[r["from_person"], r["amount"], r.get("expected_date") or ""] for r in report["receivables"]],
               [30, 18, 16], money_cols=(2,))

    # ── Cash flow categories ──
    ws = wb.create_sheet("Cash Flow")
    _table(ws, ["Category", f"Amount ({ccy})", "Share %"],
           [[c["name"], c["amount"], round(c["share"])] for c in cf["top_categories"]],
           [28, 18, 12], money_cols=(2,), pct_cols=(3,))

    # ── Currency mix ──
    if len(report["currency_mix"]) > 1:
        ws = wb.create_sheet("Currency Mix")
        _table(ws, ["Currency", f"Value ({ccy})", "Share %"],
               [[c["code"], c["value"], round(c["share"])] for c in report["currency_mix"]],
               [14, 20, 12], money_cols=(2,), pct_cols=(3,))

    # ── Transactions appendix ──
    start, end, _ = resolve_period(period_key)
    cache: dict[str, float] = {}
    stmt = select(Transaction).where(Transaction.user_id == user_id, Transaction.transaction_date <= end)
    if start is not None:
        stmt = stmt.where(Transaction.transaction_date >= start)
    stmt = stmt.order_by(Transaction.transaction_date.asc())
    txs = (await session.execute(stmt)).scalars().all()
    tx_rows = []
    for t in txs:
        f = await _factor(t.currency or ccy, ccy, cache)
        tx_rows.append([
            t.transaction_date.isoformat(),
            (t.description or "").replace("\n", " ").strip(),
            t.transaction_type,
            round(float(t.amount), 2),
            t.currency or "",
            t.category or "",
            round(float(t.amount) * f, 2),
        ])
    ws = wb.create_sheet("Transactions")
    _table(ws, ["Date", "Description", "Type", "Amount", "Currency", "Category", f"Amount ({ccy})"],
           tx_rows, [12, 40, 10, 14, 10, 16, 16], money_cols=(4, 7))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def build_transactions_csv(
    user_id: uuid.UUID, session: AsyncSession, period_key: str, ccy: str,
) -> str:
    ccy = ccy.upper()
    start, end, _ = resolve_period(period_key)
    cache: dict[str, float] = {}

    stmt = select(Transaction).where(Transaction.user_id == user_id, Transaction.transaction_date <= end)
    if start is not None:
        stmt = stmt.where(Transaction.transaction_date >= start)
    stmt = stmt.order_by(Transaction.transaction_date.asc())
    txs = (await session.execute(stmt)).scalars().all()

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Date", "Description", "Type", "Amount", "Currency", "Category", f"Amount ({ccy})"])
    for t in txs:
        f = await _factor(t.currency or ccy, ccy, cache)
        w.writerow([
            t.transaction_date.isoformat(),
            (t.description or "").replace("\n", " ").strip(),
            t.transaction_type,
            f"{float(t.amount):.2f}",
            t.currency or "",
            t.category or "",
            f"{float(t.amount) * f:.2f}",
        ])
    return buf.getvalue()
